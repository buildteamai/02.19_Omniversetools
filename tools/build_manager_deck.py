"""BuildTeamAI Manager-Level Slide Deck — shape-based visuals, no ASCII."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── Colours ──────────────────────────────────────────────────
BG_DARK    = RGBColor(0x0F, 0x0F, 0x1A)
BG_CARD    = RGBColor(0x1C, 0x1C, 0x2E)
BG_CARD_LT = RGBColor(0x24, 0x24, 0x3A)
ACCENT     = RGBColor(0x76, 0xB9, 0x00)   # NVIDIA green
ACCENT2    = RGBColor(0x00, 0xBC, 0xD4)   # cyan
ORANGE     = RGBColor(0xFF, 0x99, 0x00)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT      = RGBColor(0xBB, 0xBB, 0xCC)
DIM        = RGBColor(0x88, 0x88, 0x99)
RED_SOFT   = RGBColor(0xFF, 0x55, 0x55)
PURPLE     = RGBColor(0xAA, 0x77, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── Helpers ──────────────────────────────────────────────────

def _bg(slide, color=BG_DARK):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, l, t, w, h, fill_color, *, border=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill_color
    if border:
        s.line.color.rgb = border
        s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def _circle(slide, l, t, size, fill_color, *, border=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, size, size)
    s.fill.solid()
    s.fill.fore_color.rgb = fill_color
    if border:
        s.line.color.rgb = border
        s.line.width = Pt(2)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def _arrow_right(slide, l, t, w, h, fill_color):
    s = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill_color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _chevron(slide, l, t, w, h, fill_color):
    s = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill_color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _down_arrow(slide, l, t, w, h, fill_color):
    s = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill_color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _text(slide, l, t, w, h, txt, *, sz=18, bold=False, color=WHITE,
          align=PP_ALIGN.LEFT, font="Segoe UI", anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.text = txt
    p.font.size = Pt(sz)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    p.alignment = align
    return tb


def _multiline(slide, l, t, w, h, lines, *, sz=16, color=WHITE,
               font="Segoe UI", spacing=6, align=PP_ALIGN.LEFT, bold_first=False):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(sz)
        p.font.name = font
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.alignment = align
        if bold_first and i == 0:
            p.font.bold = True
    return tb


def _accent_line(slide, l, t, w):
    _rect(slide, l, t, w, Inches(0.04), ACCENT)


def _card(slide, l, t, w, h, title, body_lines, *, icon_color=ACCENT, title_sz=16, body_sz=13):
    """Rounded card with colored top accent bar."""
    _rect(slide, l, t, w, h, BG_CARD, radius=True)
    _rect(slide, l, t, w, Inches(0.06), icon_color)
    _text(slide, l + Inches(0.25), t + Inches(0.2), w - Inches(0.5), Inches(0.4),
          title, sz=title_sz, bold=True, color=WHITE)
    _multiline(slide, l + Inches(0.25), t + Inches(0.65), w - Inches(0.5), h - Inches(0.85),
               body_lines, sz=body_sz, color=LIGHT, spacing=4)


def _numbered_circle(slide, l, t, num, color=ACCENT):
    _circle(slide, l, t, Inches(0.5), color)
    _text(slide, l, t, Inches(0.5), Inches(0.5),
          str(num), sz=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _stat_block(slide, l, t, value, label, color=ACCENT):
    _text(slide, l, t, Inches(2.5), Inches(0.6),
          value, sz=42, bold=True, color=color, font="Segoe UI")
    _text(slide, l, t + Inches(0.6), Inches(2.5), Inches(0.4),
          label, sz=14, color=LIGHT, font="Segoe UI")


def new_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    return s


# ══════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ══════════════════════════════════════════════════════════════
s = new_slide()
# Large green accent rectangle top-left
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.2), Inches(1.8), Inches(10), Inches(1.2),
      "BuildTeamAI", sz=64, bold=True, color=WHITE)
_text(s, Inches(1.2), Inches(3.1), Inches(10), Inches(0.8),
      "Industrial Digital Twin Platform", sz=32, color=ACCENT)
_accent_line(s, Inches(1.2), Inches(3.9), Inches(3))
_multiline(s, Inches(1.2), Inches(4.2), Inches(8), Inches(1.5), [
    "Parametric 3D geometry for industrial facilities",
    "Built on NVIDIA Omniverse  •  USD-native pipeline",
    "From manufacturer specs to digital twin — no CAD required",
], sz=20, color=LIGHT, spacing=10)
_text(s, Inches(1.2), Inches(6.5), Inches(10), Inches(0.4),
      "Technical Overview  •  February 2026", sz=16, color=DIM)

# ══════════════════════════════════════════════════════════════
# SLIDE 2 — The Problem
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "The Problem", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

# Pain point cards
cards = [
    ("Heavy CAD Dependencies",
     ["Revit, SolidWorks, AutoCAD licenses",
      "per seat — expensive and siloed",
      "Files are rigid, hard to parameterize"],
     RED_SOFT),
    ("Manual Modeling Bottleneck",
     ["Each equipment piece modeled by hand",
      "Change a dimension? Start over",
      "Weeks of work for facility layouts"],
     ORANGE),
    ("Fragmented Collaboration",
     ["Structural, MEP, architecture in",
      "separate tools with separate formats",
      "No single source of truth"],
     PURPLE),
    ("No Parametric Control",
     ["Imported geometry is dead mesh",
      "Can't regenerate from specifications",
      "No metadata for BOM or analysis"],
     ACCENT2),
]

x_start = Inches(0.6)
for i, (title, lines, color) in enumerate(cards):
    x = x_start + i * Inches(3.15)
    _card(s, x, Inches(1.7), Inches(2.95), Inches(2.8), title, lines,
          icon_color=color, title_sz=15, body_sz=13)

# Bottom stat bar
_rect(s, Inches(0.6), Inches(5.0), Inches(12.1), Inches(2.0), BG_CARD, radius=True)
_stat_block(s, Inches(1.2), Inches(5.3), "$8K+", "Avg. annual CAD license cost / seat", RED_SOFT)
_stat_block(s, Inches(4.5), Inches(5.3), "4-6 wks", "Typical facility model turnaround", ORANGE)
_stat_block(s, Inches(8.0), Inches(5.3), "3-5 tools", "Required for full facility coverage", PURPLE)

# ══════════════════════════════════════════════════════════════
# SLIDE 3 — Our Solution
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Our Solution", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

_text(s, Inches(1.0), Inches(1.4), Inches(11), Inches(0.6),
      "Spec Sheet  →  Construction Algorithm  →  USD Digital Twin",
      sz=24, bold=True, color=ACCENT)

# Three-step flow with chevrons
steps = [
    ("Manufacturer\nSpecification", "JSON data, AISC tables,\ndatasheet parameters", ACCENT),
    ("Parametric\nGenerator", "build123d B-Rep kernel\ncreates solid geometry", ACCENT2),
    ("USD Scene\non Omniverse", "Meshes, materials, metadata\nready for collaboration", ORANGE),
]

for i, (title, desc, color) in enumerate(steps):
    x = Inches(1.0) + i * Inches(4.0)
    # Step card
    _rect(s, x, Inches(2.3), Inches(3.2), Inches(2.2), BG_CARD, border=color, radius=True)
    _numbered_circle(s, x + Inches(0.2), Inches(2.5), i + 1, color)
    _text(s, x + Inches(0.9), Inches(2.5), Inches(2.1), Inches(0.65),
          title, sz=18, bold=True, color=WHITE)
    _text(s, x + Inches(0.2), Inches(3.4), Inches(2.8), Inches(0.9),
          desc, sz=14, color=LIGHT)
    # Arrow between
    if i < 2:
        _arrow_right(s, x + Inches(3.3), Inches(3.2), Inches(0.6), Inches(0.4), color)

# Key benefits row
_rect(s, Inches(0.6), Inches(5.0), Inches(12.1), Inches(2.1), BG_CARD, radius=True)
benefits = [
    ("Zero CAD licenses", "Generate geometry directly\nfrom specs — no imports"),
    ("Fully parametric", "Change a number,\nregenerate instantly"),
    ("USD-native output", "Real-time collaboration\nin NVIDIA Omniverse"),
    ("Engineering-grade", "B-Rep solid modeling,\nnot mesh approximations"),
]
for i, (title, desc) in enumerate(benefits):
    x = Inches(1.0) + i * Inches(3.0)
    _text(s, x, Inches(5.2), Inches(2.5), Inches(0.35), title,
          sz=16, bold=True, color=ACCENT)
    _text(s, x, Inches(5.6), Inches(2.5), Inches(0.9), desc,
          sz=13, color=LIGHT)

# ══════════════════════════════════════════════════════════════
# SLIDE 4 — Platform Architecture (visual)
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Platform Architecture", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

# Layer diagram — stacked horizontal bars
layers = [
    ("Omniverse Kit Runtime", "Real-time 3D viewport, physics, rendering", RGBColor(0x33, 0x33, 0x55)),
    ("BuildTeamAI Extension", "UI windows, equipment catalog, tools menu", BG_CARD_LT),
    ("Object Generators", "Structural  •  MEP  •  Components", RGBColor(0x2A, 0x3A, 0x1A)),
    ("build123d CAD Kernel", "B-Rep solid modeling (OpenCascade)", RGBColor(0x1A, 0x2A, 0x3A)),
    ("USD Stage", "Meshes  •  Materials  •  Metadata  •  Mating Ports", RGBColor(0x3A, 0x2A, 0x1A)),
]

layer_w = Inches(8.5)
layer_h = Inches(0.85)
x0 = Inches(0.8)
y0 = Inches(1.5)

for i, (name, desc, color) in enumerate(layers):
    y = y0 + i * (layer_h + Inches(0.15))
    _rect(s, x0, y, layer_w, layer_h, color, radius=True)
    _text(s, x0 + Inches(0.4), y + Inches(0.08), Inches(4), Inches(0.4),
          name, sz=16, bold=True, color=WHITE)
    _text(s, x0 + Inches(0.4), y + Inches(0.42), Inches(7), Inches(0.35),
          desc, sz=12, color=LIGHT)
    # Down arrows between layers
    if i < len(layers) - 1:
        _down_arrow(s, x0 + Inches(4.0), y + layer_h, Inches(0.35), Inches(0.12), DIM)

# Side modules
side_x = Inches(9.8)
side_modules = [
    ("Import / Export", "STEP, OBJ, FBX, glTF", ACCENT2),
    ("TripoSR AI", "Image → 3D mesh", ORANGE),
    ("Frame Solver", "Structural assembly", PURPLE),
    ("Mating System", "Port-based connections", ACCENT),
]
for i, (name, desc, color) in enumerate(side_modules):
    y = Inches(1.5) + i * Inches(1.25)
    _rect(s, side_x, y, Inches(3.0), Inches(1.0), BG_CARD, border=color, radius=True)
    _text(s, side_x + Inches(0.2), y + Inches(0.1), Inches(2.5), Inches(0.35),
          name, sz=14, bold=True, color=color)
    _text(s, side_x + Inches(0.2), y + Inches(0.5), Inches(2.5), Inches(0.35),
          desc, sz=12, color=LIGHT)

# ══════════════════════════════════════════════════════════════
# SLIDE 5 — Structural Steel
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Structural Steel Library", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

# Feature cards in 2x2 grid
struct_features = [
    ("Wide Flange Beams", [
        "Full AISC W-beam catalog",
        "Data-driven from specification database",
        "Any section generated on demand",
        "Real cross-section profiles, not boxes",
    ], ACCENT),
    ("Channels & HSS Tubes", [
        "C-channels with real flange geometry",
        "Hollow structural sections (square/rect)",
        "Parametric wall thickness and dimensions",
        "Standard sizing from steel tables",
    ], ACCENT2),
    ("Frame Solver", [
        "Define members by start/end points",
        "Auto-orient along member vectors",
        "Generate connections at joints",
        "Complete frame as USD hierarchy",
    ], ORANGE),
    ("Strongbacks & Connections", [
        "Support structures for MEP systems",
        "Bolted / welded connection details",
        "Parametric plate and gusset geometry",
        "Integrated with mating port system",
    ], PURPLE),
]

for i, (title, lines, color) in enumerate(struct_features):
    col = i % 2
    row = i // 2
    x = Inches(0.8) + col * Inches(6.2)
    y = Inches(1.5) + row * Inches(2.7)
    _card(s, x, y, Inches(5.8), Inches(2.4), title, lines,
          icon_color=color, title_sz=18, body_sz=14)

_text(s, Inches(0.8), Inches(6.9), Inches(12), Inches(0.4),
      "All geometry from AISC specification data — zero pre-modeled assets required",
      sz=15, color=ORANGE, font="Segoe UI")

# ══════════════════════════════════════════════════════════════
# SLIDE 6 — MEP Systems
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "MEP Systems", sz=36, bold=True, color=WHITE)
_text(s, Inches(1.0), Inches(0.95), Inches(10), Inches(0.4),
      "Mechanical  •  Electrical  •  Plumbing", sz=18, color=DIM)
_accent_line(s, Inches(1.0), Inches(1.35), Inches(2))

mep_features = [
    ("Ductwork", [
        "Parametric duct runs with transitions",
        "Round and rectangular cross-sections",
        "Automatic sheet metal thickness",
        "Reducer/expander transitions",
    ], ACCENT),
    ("Trapeze Hangers", [
        "Support systems for ducts and pipes",
        "Configurable span, rod diameter, channel",
        "Auto-sized to carried load",
        "Mating ports for beam attachment",
    ], ACCENT2),
    ("Piping", [
        "Parametric pipe runs",
        "Standard nominal pipe sizes",
        "Elbow and tee fittings",
        "Material and schedule selection",
    ], ORANGE),
]

for i, (title, lines, color) in enumerate(mep_features):
    x = Inches(0.6) + i * Inches(4.15)
    _card(s, x, Inches(1.7), Inches(3.95), Inches(2.8), title, lines,
          icon_color=color, title_sz=18, body_sz=14)

# How it works strip
_rect(s, Inches(0.6), Inches(4.9), Inches(12.1), Inches(2.2), BG_CARD, radius=True)
_text(s, Inches(1.0), Inches(5.1), Inches(10), Inches(0.4),
      "How It Works", sz=20, bold=True, color=WHITE)

flow_steps = [
    ("Select Type", "Duct, pipe,\nor hanger"),
    ("Set Parameters", "Diameter, length,\nmaterial, gauge"),
    ("Generate", "build123d creates\nB-Rep solid"),
    ("Place in Scene", "USD mesh with\nmetadata + ports"),
]
for i, (title, desc) in enumerate(flow_steps):
    x = Inches(1.0) + i * Inches(3.0)
    _numbered_circle(s, x, Inches(5.55), i + 1, ACCENT)
    _text(s, x + Inches(0.65), Inches(5.5), Inches(2.0), Inches(0.35),
          title, sz=14, bold=True, color=WHITE)
    _text(s, x + Inches(0.65), Inches(5.9), Inches(2.0), Inches(0.7),
          desc, sz=12, color=LIGHT)
    if i < 3:
        _arrow_right(s, x + Inches(2.55), Inches(5.65), Inches(0.35), Inches(0.25), DIM)

# ══════════════════════════════════════════════════════════════
# SLIDE 7 — Components & Enclosures
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Components & Enclosures", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

comps = [
    ("Sheet Metal Panels", "Flat and formed panels\nwith bend allowances", ACCENT),
    ("Pyramids / Hoppers", "Funnel geometry with\nparametric fillets", ACCENT2),
    ("Screen Guards", "Perforated protective\nenclosures and covers", ORANGE),
    ("Stairs", "Code-compliant geometry\nrise / run / width", PURPLE),
    ("Construction Cubes", "Modular enclosure\nbuilding blocks", RGBColor(0xFF, 0x66, 0x99)),
]

for i, (title, desc, color) in enumerate(comps):
    x = Inches(0.5) + i * Inches(2.55)
    _rect(s, x, Inches(1.5), Inches(2.35), Inches(2.5), BG_CARD, border=color, radius=True)
    # Color circle icon
    _circle(s, x + Inches(0.8), Inches(1.7), Inches(0.7), color)
    _text(s, x + Inches(0.2), Inches(2.6), Inches(1.95), Inches(0.4),
          title, sz=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    _text(s, x + Inches(0.2), Inches(3.1), Inches(1.95), Inches(0.7),
          desc, sz=12, color=LIGHT, align=PP_ALIGN.CENTER)

# Parametric advantage callout
_rect(s, Inches(0.6), Inches(4.5), Inches(12.1), Inches(2.6), BG_CARD, radius=True)
_text(s, Inches(1.0), Inches(4.7), Inches(10), Inches(0.5),
      "Every Component Is Fully Parametric", sz=22, bold=True, color=ACCENT)

param_benefits = [
    ("Dimension Control", "Adjust height, width, thickness,\nangles — regenerate in real-time"),
    ("Metadata-Rich", "Every prim carries generator type\nand parameters for BOM extraction"),
    ("Material Assignment", "UsdPreviewSurface materials applied\nautomatically — steel, aluminum, etc."),
    ("Regeneration", "Change any parameter later —\ngeometry rebuilds, connections preserved"),
]
for i, (title, desc) in enumerate(param_benefits):
    x = Inches(1.0) + i * Inches(3.0)
    _text(s, x, Inches(5.3), Inches(2.5), Inches(0.35), title,
          sz=14, bold=True, color=WHITE)
    _text(s, x, Inches(5.7), Inches(2.5), Inches(0.9), desc,
          sz=12, color=LIGHT)

# ══════════════════════════════════════════════════════════════
# SLIDE 8 — Equipment Catalog (IEIP)
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Equipment Catalog", sz=36, bold=True, color=WHITE)
_text(s, Inches(1.0), Inches(0.95), Inches(10), Inches(0.4),
      "Industrial Equipment Intelligence Platform (IEIP)", sz=18, color=DIM)
_accent_line(s, Inches(1.0), Inches(1.35), Inches(2))

# Simulated catalog UI mockup using shapes
ui_x, ui_y = Inches(0.6), Inches(1.7)
ui_w, ui_h = Inches(7.5), Inches(5.3)

# Window frame
_rect(s, ui_x, ui_y, ui_w, ui_h, RGBColor(0x2A, 0x2A, 0x3E), border=DIM, radius=True)
# Title bar
_rect(s, ui_x, ui_y, ui_w, Inches(0.45), RGBColor(0x33, 0x33, 0x4A))
_text(s, ui_x + Inches(0.15), ui_y + Inches(0.05), Inches(4), Inches(0.35),
      "Equipment Catalog", sz=12, bold=True, color=LIGHT, font="Consolas")

# Left panel (tree)
tree_x = ui_x + Inches(0.1)
tree_y = ui_y + Inches(0.55)
_rect(s, tree_x, tree_y, Inches(2.3), Inches(4.6), RGBColor(0x1A, 0x1A, 0x2A), radius=True)
tree_items = [
    ("  Structural", True),
    ("    Wide Flange", False),
    ("    Channel", False),
    ("    HSS Tube", False),
    ("  MEP Systems", True),
    ("    Ductwork", False),
    ("    Trapeze", False),
    ("  Components", True),
    ("    Sheet Metal", False),
    ("    Pyramid", False),
    ("    Screen Guard", False),
    ("    Stairs", False),
]
for i, (item, is_header) in enumerate(tree_items):
    _text(s, tree_x + Inches(0.1), tree_y + Inches(0.1) + i * Inches(0.35),
          Inches(2.0), Inches(0.3), item,
          sz=10, bold=is_header, color=ACCENT if is_header else LIGHT, font="Consolas")

# Right panel (params)
param_x = tree_x + Inches(2.5)
_rect(s, param_x, tree_y, Inches(4.5), Inches(4.6), RGBColor(0x1A, 0x1A, 0x2A), radius=True)
_text(s, param_x + Inches(0.2), tree_y + Inches(0.1), Inches(3), Inches(0.3),
      "Wide Flange Parameters", sz=12, bold=True, color=ACCENT, font="Consolas")

param_fields = [
    ("Section Size:", "W12x26"),
    ("Length:", "180.0 in"),
    ("Material:", "A992 Steel"),
    ("Detail Level:", "PARAMETRIC (2)"),
]
for i, (label, value) in enumerate(param_fields):
    y = tree_y + Inches(0.6) + i * Inches(0.5)
    _text(s, param_x + Inches(0.2), y, Inches(1.5), Inches(0.3),
          label, sz=10, color=DIM, font="Consolas")
    _rect(s, param_x + Inches(1.8), y, Inches(2.2), Inches(0.32),
          RGBColor(0x22, 0x22, 0x35), radius=True)
    _text(s, param_x + Inches(1.9), y + Inches(0.02), Inches(2.0), Inches(0.28),
          value, sz=10, color=WHITE, font="Consolas")

# Place button
_rect(s, param_x + Inches(1.5), tree_y + Inches(3.0), Inches(1.8), Inches(0.45), ACCENT, radius=True)
_text(s, param_x + Inches(1.5), tree_y + Inches(3.0), Inches(1.8), Inches(0.45),
      "Place in Scene", sz=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# Right side — key features
feat_x = Inches(8.5)
_text(s, feat_x, Inches(1.7), Inches(4.5), Inches(0.5),
      "Key Capabilities", sz=20, bold=True, color=WHITE)

catalog_features = [
    ("Unified Browser", "Browse all equipment types\nin one searchable interface"),
    ("Parameter Control", "Adjust every dimension\nbefore placement"),
    ("Progressive Detail", "4 LOD levels: placeholder\nthrough fully detailed"),
    ("Spec Import", "Paste manufacturer JSON\nto create custom equipment"),
    ("Assembly Templates", "Multi-component systems\nfrom JSON definitions"),
]
for i, (title, desc) in enumerate(catalog_features):
    y = Inches(2.3) + i * Inches(0.95)
    _circle(s, feat_x, y + Inches(0.05), Inches(0.12), ACCENT)
    _text(s, feat_x + Inches(0.25), y, Inches(4.0), Inches(0.3),
          title, sz=14, bold=True, color=WHITE)
    _text(s, feat_x + Inches(0.25), y + Inches(0.3), Inches(4.0), Inches(0.55),
          desc, sz=11, color=LIGHT)

# ══════════════════════════════════════════════════════════════
# SLIDE 9 — Progressive Detail Levels
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Progressive Detail Levels", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

_text(s, Inches(1.0), Inches(1.4), Inches(11), Inches(0.5),
      "Right level of detail for every project phase — from concept to construction",
      sz=18, color=LIGHT)

levels = [
    ("Level 0", "PLACEHOLDER", "Simple bounding box\nfor space planning\nand early layout",
     RGBColor(0x44, 0x44, 0x66), DIM),
    ("Level 1", "SCHEMATIC", "Simplified outline\nfor design reviews\nand coordination",
     RGBColor(0x33, 0x44, 0x55), ACCENT2),
    ("Level 2", "PARAMETRIC", "Full geometry from\nspecifications — the\nprimary working level",
     RGBColor(0x2A, 0x3A, 0x1A), ACCENT),
    ("Level 3", "DETAILED", "Cosmetic details:\nfillets, threads, labels\nfor final deliverables",
     RGBColor(0x3A, 0x2A, 0x1A), ORANGE),
]

for i, (level, name, desc, bg_col, accent_col) in enumerate(levels):
    x = Inches(0.6) + i * Inches(3.15)
    # Card
    _rect(s, x, Inches(2.2), Inches(2.95), Inches(4.0), bg_col, border=accent_col, radius=True)
    # Level badge
    _rect(s, x + Inches(0.2), Inches(2.4), Inches(1.0), Inches(0.4), accent_col, radius=True)
    _text(s, x + Inches(0.2), Inches(2.4), Inches(1.0), Inches(0.4),
          level, sz=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # Name
    _text(s, x + Inches(0.2), Inches(3.0), Inches(2.5), Inches(0.4),
          name, sz=20, bold=True, color=accent_col)
    # Description
    _text(s, x + Inches(0.2), Inches(3.5), Inches(2.5), Inches(1.2),
          desc, sz=14, color=LIGHT)

    # Visual representation — nested rectangles showing complexity
    box_y = Inches(4.8)
    if i == 0:
        _rect(s, x + Inches(0.6), box_y, Inches(1.7), Inches(1.0), accent_col, radius=True)
    elif i == 1:
        _rect(s, x + Inches(0.5), box_y, Inches(1.9), Inches(1.0), RGBColor(0x1A, 0x1A, 0x2E), border=accent_col, radius=True)
        _rect(s, x + Inches(0.7), box_y + Inches(0.15), Inches(0.5), Inches(0.7), accent_col, radius=True)
        _rect(s, x + Inches(1.3), box_y + Inches(0.25), Inches(0.9), Inches(0.5), accent_col, radius=True)
    elif i == 2:
        _rect(s, x + Inches(0.5), box_y, Inches(1.9), Inches(1.0), RGBColor(0x1A, 0x1A, 0x2E), border=accent_col, radius=True)
        _rect(s, x + Inches(0.6), box_y + Inches(0.1), Inches(0.5), Inches(0.8), accent_col, radius=True)
        _rect(s, x + Inches(1.15), box_y + Inches(0.2), Inches(1.1), Inches(0.6), accent_col, radius=True)
        _rect(s, x + Inches(1.2), box_y + Inches(0.25), Inches(0.3), Inches(0.5), bg_col, radius=True)
    elif i == 3:
        _rect(s, x + Inches(0.5), box_y, Inches(1.9), Inches(1.0), RGBColor(0x1A, 0x1A, 0x2E), border=accent_col, radius=True)
        _rect(s, x + Inches(0.6), box_y + Inches(0.1), Inches(0.5), Inches(0.8), accent_col, radius=True)
        _rect(s, x + Inches(1.15), box_y + Inches(0.15), Inches(1.1), Inches(0.7), accent_col, radius=True)
        _rect(s, x + Inches(1.2), box_y + Inches(0.2), Inches(0.3), Inches(0.6), bg_col, radius=True)
        _circle(s, x + Inches(0.65), box_y + Inches(0.15), Inches(0.15), ORANGE)
        _circle(s, x + Inches(1.55), box_y + Inches(0.65), Inches(0.12), ORANGE)
        _circle(s, x + Inches(1.9), box_y + Inches(0.15), Inches(0.12), ORANGE)

    # Arrow between cards
    if i < 3:
        _arrow_right(s, x + Inches(3.05), Inches(3.8), Inches(0.25), Inches(0.2), DIM)

# ══════════════════════════════════════════════════════════════
# SLIDE 10 — Assembly System
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Assembly System", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

_text(s, Inches(1.0), Inches(1.35), Inches(11), Inches(0.5),
      "Multi-component systems from JSON templates — define once, place anywhere",
      sz=18, color=LIGHT)

# Flow: Template → Build → Scene
# Template card
_rect(s, Inches(0.6), Inches(2.2), Inches(3.8), Inches(4.5), BG_CARD, border=ACCENT, radius=True)
_text(s, Inches(0.8), Inches(2.35), Inches(3.4), Inches(0.4),
      "JSON Template", sz=18, bold=True, color=ACCENT)
_rect(s, Inches(0.8), Inches(2.9), Inches(3.4), Inches(3.5), RGBColor(0x16, 0x16, 0x28), radius=True)
template_lines = [
    '  "assembly_id": "trapeze_run"',
    '  "components": [',
    '    {',
    '      "entry_id": "trapeze"',
    '      "params": { "span": 48 }',
    '      "ports": ["top_left"]',
    '    },',
    '    {',
    '      "entry_id": "wide_flange"',
    '      "params": { "size": "W8x18" }',
    '      "mate_to": { port: "top_left" }',
    '    }',
    '  ]',
]
_multiline(s, Inches(0.95), Inches(3.05), Inches(3.1), Inches(3.2),
           template_lines, sz=10, color=LIGHT, font="Consolas", spacing=2)

# Arrow 1
_arrow_right(s, Inches(4.5), Inches(4.2), Inches(0.7), Inches(0.5), ACCENT)

# Build process card
_rect(s, Inches(5.3), Inches(2.2), Inches(3.2), Inches(4.5), BG_CARD, border=ACCENT2, radius=True)
_text(s, Inches(5.5), Inches(2.35), Inches(2.8), Inches(0.4),
      "Assembly Builder", sz=18, bold=True, color=ACCENT2)

build_steps = [
    "1. Look up each component\n   in equipment registry",
    "2. Invoke generator with\n   specified parameters",
    "3. Resolve mating ports —\n   align positions + normals",
    "4. Group all components\n   under parent Xform prim",
    "5. Write assembly metadata\n   to USD custom data",
]
_multiline(s, Inches(5.5), Inches(3.0), Inches(2.8), Inches(3.5),
           build_steps, sz=12, color=LIGHT, spacing=8)

# Arrow 2
_arrow_right(s, Inches(8.6), Inches(4.2), Inches(0.7), Inches(0.5), ACCENT2)

# Result card
_rect(s, Inches(9.4), Inches(2.2), Inches(3.5), Inches(4.5), BG_CARD, border=ORANGE, radius=True)
_text(s, Inches(9.6), Inches(2.35), Inches(3.1), Inches(0.4),
      "USD Scene Result", sz=18, bold=True, color=ORANGE)

result_items = [
    "/World/TrapezRun  (Xform)",
    "  /Trapeze_001  (Mesh)",
    "    twin:generator = trapeze",
    "    twin:param:span = 48",
    "    twin:is_port = top_left",
    "",
    "  /WFlange_001  (Mesh)",
    "    twin:generator = wide_flange",
    "    twin:param:size = W8x18",
    "    twin:mated_to = top_left",
    "",
    "  assembly_metadata",
    "    ieip:assembly_id = trapeze_run",
]
_multiline(s, Inches(9.6), Inches(3.0), Inches(3.1), Inches(3.5),
           result_items, sz=10, color=LIGHT, font="Consolas", spacing=2)

# ══════════════════════════════════════════════════════════════
# SLIDE 11 — TripoSR Image-to-3D
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "TripoSR — Image to 3D", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))
_text(s, Inches(1.0), Inches(1.35), Inches(11), Inches(0.5),
      "Upload a photo of any equipment — get a 3D mesh in ~10 seconds",
      sz=18, color=LIGHT)

# Pipeline flow: Photo → Server → Mesh → Scene
pipe_items = [
    ("Photo Input", "Upload any\nequipment image", ACCENT, "IMAGE"),
    ("TripoSR Server", "FastAPI on localhost\nPyTorch + CUDA", ACCENT2, "AI"),
    ("3D Mesh", "OBJ mesh output\n~5.8 MB typical", ORANGE, "OBJ"),
    ("USD Scene", "Imported via\nmesh_importer.py", PURPLE, "USD"),
]

for i, (title, desc, color, badge) in enumerate(pipe_items):
    x = Inches(0.6) + i * Inches(3.2)
    _rect(s, x, Inches(2.2), Inches(2.8), Inches(2.3), BG_CARD, border=color, radius=True)
    # Badge
    _rect(s, x + Inches(0.2), Inches(2.4), Inches(0.7), Inches(0.35), color, radius=True)
    _text(s, x + Inches(0.2), Inches(2.4), Inches(0.7), Inches(0.35),
          badge, sz=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _text(s, x + Inches(0.2), Inches(2.95), Inches(2.4), Inches(0.4),
          title, sz=16, bold=True, color=WHITE)
    _text(s, x + Inches(0.2), Inches(3.4), Inches(2.4), Inches(0.8),
          desc, sz=13, color=LIGHT)
    if i < 3:
        _arrow_right(s, x + Inches(2.9), Inches(3.2), Inches(0.25), Inches(0.2), DIM)

# Performance stats
_rect(s, Inches(0.6), Inches(4.9), Inches(12.1), Inches(2.2), BG_CARD, radius=True)
_text(s, Inches(1.0), Inches(5.1), Inches(10), Inches(0.4),
      "Performance", sz=20, bold=True, color=WHITE)

perf_stats = [
    ("~10s", "Total pipeline", ACCENT),
    ("5.4s", "Model load", ACCENT2),
    ("1.7s", "Inference", ORANGE),
    ("2.4s", "Mesh extraction", PURPLE),
]
for i, (val, label, color) in enumerate(perf_stats):
    x = Inches(1.0) + i * Inches(2.5)
    _text(s, x, Inches(5.6), Inches(2.0), Inches(0.6), val,
          sz=36, bold=True, color=color)
    _text(s, x, Inches(6.15), Inches(2.0), Inches(0.3), label,
          sz=13, color=LIGHT)

_text(s, Inches(8.0), Inches(5.6), Inches(4.5), Inches(0.3),
      "Hardware Requirements", sz=14, bold=True, color=WHITE)
_multiline(s, Inches(8.0), Inches(5.95), Inches(4.5), Inches(1.0), [
    "NVIDIA RTX GPU (12 GB+ VRAM)",
    "PyTorch 2.5.1 + CUDA 12.1",
    "Model weights: ~1.68 GB (one-time download)",
], sz=12, color=LIGHT, spacing=3)

# ══════════════════════════════════════════════════════════════
# SLIDE 12 — Import / Export
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Import & Export", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

# Import side
_rect(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(5.2), BG_CARD, radius=True)
_text(s, Inches(1.0), Inches(1.8), Inches(5), Inches(0.5),
      "Import Pipelines", sz=22, bold=True, color=ACCENT)

import_items = [
    ("STEP / STP", "Engineering CAD files via OpenCascade\nPreserves assembly structure as USD hierarchy", ACCENT),
    ("OBJ / FBX / glTF", "Mesh files via mesh_importer.py\nVertex + face extraction to UsdGeom.Mesh", ACCENT2),
    ("Image → 3D", "Any photo via TripoSR AI\nGenerates OBJ mesh, imports to USD", ORANGE),
]

for i, (fmt, desc, color) in enumerate(import_items):
    y = Inches(2.5) + i * Inches(1.3)
    _rect(s, Inches(0.9), y, Inches(1.3), Inches(0.4), color, radius=True)
    _text(s, Inches(0.9), y, Inches(1.3), Inches(0.4), fmt,
          sz=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _arrow_right(s, Inches(2.3), y + Inches(0.05), Inches(0.4), Inches(0.3), color)
    _text(s, Inches(2.9), y - Inches(0.05), Inches(3.3), Inches(0.7), desc,
          sz=12, color=LIGHT)

# Export side
_rect(s, Inches(6.8), Inches(1.6), Inches(5.9), Inches(5.2), BG_CARD, radius=True)
_text(s, Inches(7.2), Inches(1.8), Inches(5), Inches(0.5),
      "Export Pipelines", sz=22, bold=True, color=ORANGE)

export_items = [
    ("STEP Export", "Re-invoke generator with stored metadata\nbuild123d solid → OCC STEP writer → .STEP", ORANGE),
    ("USD Native", "Standard Omniverse output\nShareable via Nucleus server", ACCENT),
    ("BOM Extract", "Bill of Materials from prim metadata\nJSON output for procurement systems", PURPLE),
]

for i, (fmt, desc, color) in enumerate(export_items):
    y = Inches(2.5) + i * Inches(1.3)
    _rect(s, Inches(7.1), y, Inches(1.5), Inches(0.4), color, radius=True)
    _text(s, Inches(7.1), y, Inches(1.5), Inches(0.4), fmt,
          sz=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _arrow_right(s, Inches(8.7), y + Inches(0.05), Inches(0.4), Inches(0.3), color)
    _text(s, Inches(9.3), y - Inches(0.05), Inches(3.1), Inches(0.7), desc,
          sz=12, color=LIGHT)

# Round-trip callout
_rect(s, Inches(0.6), Inches(6.0), Inches(12.1), Inches(0.7), RGBColor(0x2A, 0x3A, 0x1A), radius=True)
_text(s, Inches(1.0), Inches(6.05), Inches(11), Inches(0.6),
      "Full round-trip:  STEP → USD (edit parametrically) → STEP   — no data loss, metadata preserved",
      sz=15, bold=True, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)

# ══════════════════════════════════════════════════════════════
# SLIDE 13 — Mating & Connections
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Mating & Connections", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))
_text(s, Inches(1.0), Inches(1.35), Inches(11), Inches(0.5),
      "Snap components together at defined anchor points — connections survive regeneration",
      sz=18, color=LIGHT)

# Visual: Two components with port dots connecting
# Component A
_rect(s, Inches(1.0), Inches(2.5), Inches(4.0), Inches(2.5), BG_CARD, border=ACCENT, radius=True)
_text(s, Inches(1.3), Inches(2.7), Inches(3.4), Inches(0.4),
      "Component A — Trapeze Hanger", sz=14, bold=True, color=ACCENT)
# Simulated geometry
_rect(s, Inches(1.5), Inches(3.3), Inches(3.0), Inches(0.15), LIGHT)  # channel
_rect(s, Inches(1.8), Inches(3.2), Inches(0.08), Inches(1.2), LIGHT)  # left rod
_rect(s, Inches(4.1), Inches(3.2), Inches(0.08), Inches(1.2), LIGHT)  # right rod
# Port dots
_circle(s, Inches(1.7), Inches(3.15), Inches(0.2), ORANGE, border=WHITE)
_text(s, Inches(1.4), Inches(3.4), Inches(1.0), Inches(0.3),
      "top_left", sz=9, color=ORANGE, font="Consolas")
_circle(s, Inches(4.05), Inches(3.15), Inches(0.2), ORANGE, border=WHITE)
_text(s, Inches(3.7), Inches(3.4), Inches(1.0), Inches(0.3),
      "top_right", sz=9, color=ORANGE, font="Consolas")

# Connection lines
_rect(s, Inches(5.2), Inches(3.25), Inches(1.5), Inches(0.03), ORANGE)
_rect(s, Inches(5.2), Inches(3.5), Inches(1.5), Inches(0.03), ORANGE)
_text(s, Inches(5.3), Inches(3.55), Inches(1.3), Inches(0.3),
      "mate ports", sz=10, color=ORANGE, align=PP_ALIGN.CENTER)

# Component B
_rect(s, Inches(6.9), Inches(2.5), Inches(5.5), Inches(2.5), BG_CARD, border=ACCENT2, radius=True)
_text(s, Inches(7.2), Inches(2.7), Inches(5.0), Inches(0.4),
      "Component B — W-Beam", sz=14, bold=True, color=ACCENT2)
# Simulated I-beam cross section appearance
_rect(s, Inches(7.5), Inches(3.3), Inches(4.5), Inches(0.12), LIGHT)  # top flange
_rect(s, Inches(9.5), Inches(3.3), Inches(0.08), Inches(1.0), LIGHT)  # web
_rect(s, Inches(7.5), Inches(4.2), Inches(4.5), Inches(0.12), LIGHT)  # bottom flange
_circle(s, Inches(7.35), Inches(3.2), Inches(0.2), ORANGE, border=WHITE)

# How mating works
_rect(s, Inches(0.6), Inches(5.4), Inches(12.1), Inches(1.7), BG_CARD, radius=True)
_text(s, Inches(1.0), Inches(5.55), Inches(10), Inches(0.4),
      "How Port-Based Mating Works", sz=18, bold=True, color=WHITE)

mate_steps = [
    ("Define Ports", "Generators declare anchor\npoints with position + normal"),
    ("Select Mate", "User picks source port\non component A"),
    ("Align", "System computes transform\nto align port positions"),
    ("Connect", "Metadata records connection\nsurvives regen / edit"),
]
for i, (title, desc) in enumerate(mate_steps):
    x = Inches(1.0) + i * Inches(3.0)
    _numbered_circle(s, x, Inches(5.95), i + 1, ACCENT)
    _text(s, x + Inches(0.6), Inches(5.9), Inches(2.2), Inches(0.3),
          title, sz=13, bold=True, color=WHITE)
    _text(s, x + Inches(0.6), Inches(6.2), Inches(2.2), Inches(0.65),
          desc, sz=11, color=LIGHT)

# ══════════════════════════════════════════════════════════════
# SLIDE 14 — Workflow Example
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "End-to-End Workflow", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

steps = [
    ("Open Catalog", "Launch Equipment Catalog\nfrom Tools menu", ACCENT),
    ("Browse & Configure", "Select equipment type,\nadjust parameters", ACCENT2),
    ("Place in Scene", "One-click placement\nas USD prim", ORANGE),
    ("Mate Components", "Snap structural + MEP\nat anchor ports", PURPLE),
    ("Apply Materials", "Assign materials via\nStyle Editor", RGBColor(0xFF, 0x66, 0x99)),
    ("Export", "Extract BOM, export\nSTEP or share USD", ACCENT),
]

# Horizontal timeline
_rect(s, Inches(1.5), Inches(3.0), Inches(10.5), Inches(0.04), DIM)

for i, (title, desc, color) in enumerate(steps):
    x = Inches(0.8) + i * Inches(2.05)
    # Circle on timeline
    _circle(s, x + Inches(0.65), Inches(2.78), Inches(0.45), color)
    _text(s, x + Inches(0.65), Inches(2.78), Inches(0.45), Inches(0.45),
          str(i + 1), sz=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # Card below
    _rect(s, x, Inches(3.5), Inches(1.9), Inches(2.2), BG_CARD, border=color, radius=True)
    _text(s, x + Inches(0.15), Inches(3.65), Inches(1.6), Inches(0.4),
          title, sz=14, bold=True, color=color)
    _text(s, x + Inches(0.15), Inches(4.1), Inches(1.6), Inches(0.9),
          desc, sz=12, color=LIGHT)

# Bottom callout
_rect(s, Inches(0.6), Inches(6.1), Inches(12.1), Inches(1.0), RGBColor(0x2A, 0x3A, 0x1A), radius=True)
_text(s, Inches(1.0), Inches(6.2), Inches(11), Inches(0.8),
      "Complete facility layout from manufacturer specifications — no CAD software, no manual modeling.\n"
      "Every object is parametric, metadata-rich, and ready for BOM extraction.",
      sz=15, color=LIGHT)

# ══════════════════════════════════════════════════════════════
# SLIDE 15 — Tech Stack (clean table visual)
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Technology Stack", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

stack = [
    ("Runtime",        "NVIDIA Omniverse Kit 106+",         ACCENT),
    ("Scene Format",   "USD (Universal Scene Description)",  ACCENT2),
    ("CAD Kernel",     "build123d  (OpenCascade B-Rep)",     ORANGE),
    ("GPU Compute",    "PyTorch 2.5.1 + CUDA 12.1",        PURPLE),
    ("AI Mesh Gen",    "TripoSR (Image → 3D)",             RGBColor(0xFF, 0x66, 0x99)),
    ("UI Framework",   "omni.ui",                           ACCENT),
    ("Language",       "Python 3.10+",                      ACCENT2),
    ("Data Formats",   "JSON (AISC, assemblies, specs)",    ORANGE),
    ("Import / Export","STEP, OBJ, FBX, glTF",             PURPLE),
    ("Hardware",       "NVIDIA RTX A3000 (12 GB VRAM)",    RGBColor(0xFF, 0x66, 0x99)),
]

for i, (layer, tech, color) in enumerate(stack):
    y = Inches(1.5) + i * Inches(0.55)
    # Alternating row backgrounds
    if i % 2 == 0:
        _rect(s, Inches(0.8), y, Inches(11.5), Inches(0.5), BG_CARD)
    # Color dot
    _circle(s, Inches(1.0), y + Inches(0.12), Inches(0.22), color)
    # Layer name
    _text(s, Inches(1.4), y + Inches(0.05), Inches(2.5), Inches(0.4),
          layer, sz=15, bold=True, color=WHITE)
    # Tech value
    _text(s, Inches(4.0), y + Inches(0.05), Inches(8), Inches(0.4),
          tech, sz=15, color=LIGHT)

# Footer
_text(s, Inches(0.8), Inches(7.0), Inches(12), Inches(0.3),
      "All units: inches  •  Coordinate system: Y-Up  •  All geometry: B-Rep solid modeling",
      sz=13, color=DIM)

# ══════════════════════════════════════════════════════════════
# SLIDE 16 — What's Next
# ══════════════════════════════════════════════════════════════
s = new_slide()
_rect(s, Inches(0), Inches(0), Inches(0.15), Inches(7.5), ACCENT)
_text(s, Inches(1.0), Inches(0.4), Inches(10), Inches(0.7),
      "Roadmap", sz=36, bold=True, color=WHITE)
_accent_line(s, Inches(1.0), Inches(1.1), Inches(2))

roadmap = [
    ("Now", "Current Release", [
        "Structural steel library (AISC)",
        "MEP ductwork & trapeze hangers",
        "Equipment Catalog with IEIP registry",
        "STEP import / export pipeline",
        "Port-based mating system",
    ], ACCENT),
    ("Next", "In Development", [
        "TripoSR server integration (FastAPI → Kit)",
        "Expanded equipment from manufacturer specs",
        "Assembly template library",
        "Enhanced BOM extraction and reporting",
    ], ORANGE),
    ("Future", "Planned", [
        "Multi-user collaboration via Nucleus",
        "Automated layout from P&ID drawings",
        "Physics-based clash detection",
        "Cloud rendering and review workflows",
    ], PURPLE),
]

for i, (phase, title, items, color) in enumerate(roadmap):
    x = Inches(0.6) + i * Inches(4.2)
    _rect(s, x, Inches(1.5), Inches(3.9), Inches(5.2), BG_CARD, border=color, radius=True)
    # Phase badge
    _rect(s, x + Inches(0.2), Inches(1.7), Inches(1.0), Inches(0.4), color, radius=True)
    _text(s, x + Inches(0.2), Inches(1.7), Inches(1.0), Inches(0.4), phase,
          sz=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    # Title
    _text(s, x + Inches(1.4), Inches(1.75), Inches(2.3), Inches(0.4),
          title, sz=16, bold=True, color=WHITE)
    # Items
    for j, item in enumerate(items):
        y = Inches(2.4) + j * Inches(0.65)
        _circle(s, x + Inches(0.3), y + Inches(0.05), Inches(0.12), color)
        _text(s, x + Inches(0.55), y, Inches(3.0), Inches(0.55),
              item, sz=13, color=LIGHT)

# ══════════════════════════════════════════════════════════════
# Save
# ══════════════════════════════════════════════════════════════
out_path = r"C:\programming\buildteamai\BuildTeamAI_Overview.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
print(f"Slides: {len(prs.slides)}")
